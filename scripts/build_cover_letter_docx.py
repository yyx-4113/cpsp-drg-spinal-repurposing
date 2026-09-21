#!/usr/bin/env python3
"""Build submission_pack/Cover_Letter.docx from reports/MVP_ScientificReports_cover_letter.md.
Plain-but-formatted: **bold** and *italic* inline runs parsed; Title style for the H1.
"""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(r"D:/2026.9/极速交付9月会员日优惠套路/01_AI生信-虚拟多重筛药/慢性疼痛")
SRC = ROOT / "reports" / "MVP_ScientificReports_cover_letter.md"
OUT = ROOT / "submission_pack" / "Cover_Letter.docx"


def add_inline(text: str, para):
    """Parse **bold** and *italic* into runs."""
    for bpart in re.split(r'(\*\*.+?\*\*)', text):
        if len(bpart) >= 4 and bpart.startswith('**') and bpart.endswith('**'):
            bold_text = bpart[2:-2]
            for ip in re.split(r'(\*.+?\*)', bold_text):
                if len(ip) >= 2 and ip.startswith('*') and ip.endswith('*'):
                    r = para.add_run(ip[1:-1]); r.bold = True; r.italic = True
                elif ip:
                    r = para.add_run(ip); r.bold = True
        else:
            for ip in re.split(r'(\*.+?\*)', bpart):
                if len(ip) >= 2 and ip.startswith('*') and ip.endswith('*') and not (ip.startswith('**')):
                    r = para.add_run(ip[1:-1]); r.italic = True
                elif ip:
                    para.add_run(ip)


def main():
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)

    blocks = SRC.read_text(encoding='utf-8').split('\n')
    buf = []
    for line in blocks:
        if line.startswith('# '):
            # flush any buffered paragraph
            if buf:
                p = doc.add_paragraph(); add_inline(' '.join(buf).strip(), p); buf = []
            h = doc.add_paragraph(); h.style = doc.styles['Title']
            add_inline(line[2:].strip(), h)
        elif line.strip() == '':
            if buf:
                p = doc.add_paragraph(); add_inline(' '.join(buf).strip(), p); buf = []
        else:
            buf.append(line.strip())
    if buf:
        p = doc.add_paragraph(); add_inline(' '.join(buf).strip(), p)

    doc.save(OUT)
    print(f"wrote {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
